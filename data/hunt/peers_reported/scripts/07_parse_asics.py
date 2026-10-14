"""ASICS: (a) quarterly regional sales from IR fact-sheet xlsx (FY2021, FY2022, FY2024 books; FY2023 book duplicates),
(b) 2025-2026 regional YTD sales from Japan-GAAP summary statements (tanshin) PDFs -> discrete quarters by differencing.
Currency-neutral YTD growth is captured where printed."""
import openpyxl, re, csv, os, glob
B='/home/user/nike-research-case-comp/data/hunt/peers_reported'
REGS=['Japan','North America','EMEA','Greater China','Oceania','Southeast and South Asia','Others']
rows=[]
def norm(s):
    s=str(s).split('\n')[-1].strip()
    for r in REGS+['Europe']:
        if s.startswith(r) or s.endswith(r): return 'EMEA' if r=='Europe' else r
    return None
for f,years in [('asics_factsheet_FY2021.xlsx',(2020,2021)),('asics_factsheet_FY2022.xlsx',(2021,2022)),('asics_factsheet_FY2024.xlsx',(2023,2024))]:
    wb=openpyxl.load_workbook(f'{B}/raw/asics/{f}',data_only=True)
    if f.endswith('FY2024.xlsx'):
        ws=wb.worksheets[1]; reg=None
        for r in ws.iter_rows(values_only=True):
            v=[x for x in r if x is not None]
            if not v: continue
            n=norm(v[0]) if len(v)==1 else None
            if n: reg=n; continue
            if reg and str(v[0]).endswith('Net Sales'):
                q=[v[1],v[2],v[4],v[5],v[8],v[9],v[11],v[12]]
                for k,val in enumerate(q):
                    rows.append(dict(company='asics',year=years[k//4],quarter=k%4+1,metric='net_sales_region',region=reg,value=val,yoy_cn_pct=None,unit='JPY m',source=f))
                reg=None
        ws=wb.worksheets[0]
        for r in ws.iter_rows(values_only=True):
            v=[x for x in r if x is not None]
            if v and str(v[0]).endswith(('Net Sales','Gross Profit')) and len(v)>=15:
                q=[v[1],v[2],v[4],v[5],v[8],v[9],v[11],v[12]]
                for k,val in enumerate(q):
                    rows.append(dict(company='asics',year=years[k//4],quarter=k%4+1,metric='net_sales' if 'Net' in str(v[0]) else 'gross_profit',region='Total',value=val,yoy_cn_pct=None,unit='JPY m',source=f))
                if 'Gross' in str(v[0]): break
    else:
        ws=wb.worksheets[0]; mode=None
        for r in ws.iter_rows(values_only=True):
            v=[x for x in r if x is not None]
            if not v: continue
            if v[0]=='Q1' and len(v)==8: mode='q'; continue
            if mode=='q' and str(v[0])=='Net sales' and len(v)==9:
                for k,val in enumerate(v[1:]): rows.append(dict(company='asics',year=years[k//4],quarter=k%4+1,metric='net_sales',region='Total',value=val,yoy_cn_pct=None,unit='JPY m',source=f))
            if mode=='q' and len(v)==10 and v[1]=='Net sales' and norm(v[0]):
                for k,val in enumerate(v[2:]): rows.append(dict(company='asics',year=years[k//4],quarter=k%4+1,metric='net_sales_region',region=norm(v[0]),value=val,yoy_cn_pct=None,unit='JPY m',source=f))
# tanshin YTD
T=B+'/raw/asics_text'
ytd=[]
for fn in sorted(glob.glob(T+'/asics_*.txt')):
    txt=open(fn,errors='ignore').read()
    m=re.search(r'(First Quarter|Second Quarter|Third Quarter|Fiscal Year)[^\n]*?(March 31|June 30|September 30|December 31),\s*(20\d\d)',txt)
    if not m or 'Reportable' not in txt: continue
    n={'March 31':1,'June 30':2,'September 30':3,'December 31':4}[m.group(2)]; y=int(m.group(3))
    seg=txt[txt.index('Reportable'):]
    lines=seg.splitlines()[:60]
    vals=[]
    for l in lines:
        mm=re.search(r'Net sales\s+([\d,]+)\s+([\d,]+)\s+\(?([\d,]+)\)?\s+\(?([\d\.]+)\)?\s*(\(?[\d\.]+\)?|－|-)?',l)
        if mm: vals.append(mm)
    if len(vals)<7: print('short',fn,len(vals)); continue
    regs=['Japan','North America','EMEA','Greater China','Oceania','Southeast and South Asia','Others']
    for r,mm in zip(regs,vals[:7]):
        cn=mm.group(5)
        cnv=None if (not cn or cn in '－-') else (-float(cn.strip('()')) if cn.startswith('(') else float(cn))
        ytd.append(dict(year=y,n=n,region=r,cur=float(mm.group(2).replace(',','')),prior=float(mm.group(1).replace(',','')),cn=cnv,src=os.path.basename(fn)))
    print(y,n,os.path.basename(fn),[(v['region'][:5],v['cur'],v['cn']) for v in ytd[-7:]][:3])
# write YTD table
with open(B+'/asics_ytd_regional_tanshin.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(ytd[0])); w.writeheader(); w.writerows(ytd)
# discrete quarters for 2025-2026 from YTD (current-year columns), prior-year quarter from prior columns of same doc chain
idx={(d['year'],d['n'],d['region']):d for d in ytd}
for (y,n,r),d in idx.items():
    if y<2025: continue
    if n==1: cur,pri=d['cur'],d['prior']
    else:
        p=idx.get((y,n-1,r))
        if not p: continue
        cur,pri=d['cur']-p['cur'],d['prior']-p['prior']
    rows.append(dict(company='asics',year=y,quarter=n,metric='net_sales_region',region=r,value=cur,yoy_cn_pct=None,unit='JPY m',source=d['src']+' (YTD differenced)'))
    rows.append(dict(company='asics',year=y-1,quarter=n,metric='net_sales_region_prior_restated',region=r,value=pri,yoy_cn_pct=None,unit='JPY m',source=d['src']+' (YTD differenced)'))
    if n==1 or True:
        rows.append(dict(company='asics',year=y,quarter=n,metric='ytd_cn_growth_region',region=r,value=d['cn'],yoy_cn_pct=d['cn'],unit='% YTD',source=d['src']))
with open(B+'/asics_quarterly_extract.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
import collections; c=collections.Counter((r['year'],r['quarter'],r['metric']) for r in rows)
for k in sorted(c): print(k,c[k])
