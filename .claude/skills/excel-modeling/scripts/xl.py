#!/usr/bin/env python3
"""xl: map, audit, patch, verify, diff and trace an .xlsx without rebuilding it.

    python3 xl.py map    model.xlsx [--sheet NAME] [--max-rows 80]
    python3 xl.py audit  model.xlsx [--sheet NAME] [--allow 0,1,12,100]
    python3 xl.py patch  model.xlsx ops.json --out new.xlsx [--refresh-cache]
    python3 xl.py verify model.xlsx [--cells "Sheet!B5,Sheet!C9"]
    python3 xl.py diff   old.xlsx new.xlsx [--expect "Sheet!B5:F9,Other!A1"]
    python3 xl.py trace  model.xlsx "Sheet!F12" [--depth 3] [--dependents]
    python3 xl.py features model.xlsx
    python3 xl.py scenarios model.xlsx --switch "Inputs!C4=Bear,Base,Bull" --outputs "Valuation!C10" --check "Checks!D7" --order "Valuation!C10:asc"
    python3 xl.py tornado model.xlsx "Valuation!C10" --inputs "Inputs!C17; Inputs!C14=0.2,0.25" [--pct 0.1]
    python3 xl.py render model.xlsx --out pngdir        (look at charts and layout)
    python3 xl.py sensitivity model.xlsx "Model!F9" --row "Inputs!B3=0,0.05,0.1" [--col "Inputs!B4=0.5,0.6"] [--write "Out!B2" --out new.xlsx]

Design rules (see SKILL.md): the delivered file is edited at the XML level so untouched
parts survive byte for byte; LibreOffice only ever recalculates a throwaway copy that is
used to read values and catch errors.
"""
import argparse
import json
import re
import sys
import warnings
from collections import Counter, defaultdict
from pathlib import Path

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlcore as X  # noqa: E402
from lxml import etree  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.formula.translate import Translator  # noqa: E402
from openpyxl.utils import column_index_from_string as col2i, get_column_letter as i2col  # noqa: E402
from openpyxl.utils.cell import range_boundaries  # noqa: E402

NS, M = X.NS, X.M
ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
VOLATILE = ("OFFSET", "INDIRECT", "TODAY", "NOW", "RAND", "RANDBETWEEN", "CELL", "INFO")
# Functions that LibreOffice 24.x and/or Google Sheets conversion handle badly when written
# from Python. Prefixed forms are what Excel stores; bare forms become #NAME? in Excel.
NEEDS_PREFIX = ("XLOOKUP", "XMATCH", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "LET",
                "LAMBDA", "TEXTJOIN", "CONCAT", "IFS", "SWITCH", "MAXIFS", "MINIFS", "STOCKHISTORY",
                "VSTACK", "HSTACK", "TAKE", "DROP", "CHOOSECOLS", "CHOOSEROWS", "TEXTSPLIT")
YEARLIKE = re.compile(r"^(FY|CY)?'?\d{2,4}[AEBF]?$|^(Q[1-4]|H[12])\s?(FY)?'?\d{2,4}[AEBF]?$", re.I)


# ============================================================================ helpers

def load_pair(path):
    """Formulas and cached values need two loads; never save either object."""
    wf = load_workbook(path, data_only=False)
    wv = load_workbook(path, data_only=True)
    return wf, wv


def parse_cellspec(spec, default_sheet=None):
    if "!" in spec:
        sh, addr = X.split_sheet(spec)
    else:
        sh, addr = default_sheet, spec
    return sh, addr.replace("$", "")


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def font_rgb(cell):
    try:
        c = cell.font.color
        if c is None:
            return None
        if c.type == "rgb" and isinstance(c.rgb, str):
            return c.rgb[-6:].upper()
        if c.type == "theme":
            return f"theme{c.theme}"
    except Exception:
        return None
    return None


def label_of(ws, r, max_col=3):
    for c in range(1, max_col + 1):
        v = ws.cell(r, c).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()
    return ""


def header_rows(ws, scan=15):
    """Rows that look like a timeline (3+ year/period labels)."""
    out = []
    for r in range(1, min(ws.max_row, scan) + 1):
        hits = 0
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)) and 1990 <= v <= 2100 and float(v).is_integer():
                hits += 1
            elif isinstance(v, str) and YEARLIKE.match(v.strip()):
                hits += 1
        if hits >= 3:
            out.append(r)
    return out


# ============================================================================ features

FEATURE_PATTERNS = [
    ("vba_macros", lambda n, x: n.endswith("vbaProject.bin"), "save as .xlsm and never through openpyxl without keep_vba=True"),
    ("pivot_tables", lambda n, x: "pivotTable" in n, "openpyxl keeps pivots read-only; LibreOffice rewrites them"),
    ("slicers", lambda n, x: "slicer" in n.lower(), "lost by openpyxl"),
    ("sparklines", lambda n, x: n.startswith("xl/worksheets/sheet") and b"sparklineGroup" in x, "lost by openpyxl (verified)"),
    ("dynamic_arrays", lambda n, x: n == "xl/metadata.xml", "openpyxl drops metadata.xml: spills become legacy CSE arrays (verified)"),
    ("data_tables", lambda n, x: n.startswith("xl/worksheets/sheet") and b't="dataTable"' in x, "What-If tables: keep XML-level edits only"),
    ("external_links", lambda n, x: n.startswith("xl/externalLinks/"), "cached values are the only copy of linked data"),
    ("threaded_comments", lambda n, x: "threadedComment" in n, "openpyxl keeps only legacy notes"),
    ("charts", lambda n, x: n.startswith("xl/charts/chart"), "openpyxl re-serialises charts (formatting may change)"),
    ("images", lambda n, x: n.startswith("xl/media/"), ""),
    ("tables", lambda n, x: n.startswith("xl/tables/"), "structured refs: LibreOffice returns #N/A for [@Col] (verified)"),
    ("calc_chain", lambda n, x: n == "xl/calcChain.xml", "drop after formula edits (xl.py patch does this)"),
    ("custom_xml", lambda n, x: n.startswith("customXml/"), ""),
    ("form_controls", lambda n, x: "ctrlProps" in n or "activeX" in n, "lost by openpyxl"),
    ("x14_extensions", lambda n, x: n.startswith("xl/worksheets/sheet") and b"<extLst" in x, "x14 CF/DV extensions are lost by openpyxl"),
]


def features(path):
    pkg = X.Package(path)
    found = {}
    for n in pkg.names():
        for key, test, note in FEATURE_PATTERNS:
            try:
                if test(n, pkg.raw[n]):
                    found.setdefault(key, {"parts": 0, "note": note})["parts"] += 1
            except Exception:
                pass
    wb = pkg.xml("xl/workbook.xml")
    cp = wb.find("m:calcPr", NS)
    calc = dict(cp.attrib) if cp is not None else {}
    names = []
    dn = wb.find("m:definedNames", NS)
    if dn is not None:
        for d in dn:
            names.append({"name": d.get("name"), "ref": d.text, "local": d.get("localSheetId")})
    shared = 0
    for _, part, _ in pkg.sheets():
        shared += pkg.raw[part].count(b't="shared"')
    return {"features": found, "calcPr": calc, "defined_names": names,
            "shared_formula_cells": shared,
            "safe_engine": "xml" if set(found) & {"sparklines", "dynamic_arrays", "slicers", "x14_extensions",
                                                  "form_controls", "threaded_comments", "data_tables",
                                                  "vba_macros", "pivot_tables", "external_links"} else "xml or openpyxl"}


# ============================================================================ map

def cmd_map(a):
    wf = load_workbook(a.file)
    if a.no_recalc:
        wv = load_workbook(a.file, data_only=True)
        print("(values are the file's cached values; they may be stale or blank)")
    else:
        wv = load_workbook(X.lo_recalc_copy(a.file), data_only=True)
    feat = features(a.file)
    print(f"# {a.file}")
    print(f"sheets: {', '.join(f'{ws.title}({ws.sheet_state})' for ws in wf.worksheets)}")
    print(f"calcPr: {feat['calcPr']}")
    if feat["features"]:
        print("features: " + "; ".join(f"{k} x{v['parts']}" for k, v in feat["features"].items()))
    print(f"recommended edit engine: {feat['safe_engine']}")
    if feat["defined_names"]:
        print("defined names: " + "; ".join(f"{d['name']}={d['ref']}" for d in feat["defined_names"][:40]))
    for ws in wf.worksheets:
        if a.sheet and ws.title != a.sheet:
            continue
        wsv = wv[ws.title]
        nf = nc = nt = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                if is_formula(c.value) or type(c.value).__name__ == "ArrayFormula":
                    nf += 1
                elif isinstance(c.value, (int, float)):
                    nc += 1
                else:
                    nt += 1
        hdr = header_rows(ws)
        print(f"\n## {ws.title}  dims={ws.dimensions} formulas={nf} numbers={nc} text={nt} "
              f"freeze={ws.freeze_panes} merged={len(ws.merged_cells.ranges)} "
              f"dv={len(ws.data_validations.dataValidation)} cf={len(ws.conditional_formatting)} "
              f"tables={list(ws.tables.keys())}")
        if hdr:
            r = hdr[0]
            cols = [(c, ws.cell(r, c).value) for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
            print(f"timeline row {r}: " + ", ".join(f"{i2col(c)}={v}" for c, v in cols[:30]))
        shown = 0
        for r in range(1, ws.max_row + 1):
            cells = [ws.cell(r, c) for c in range(1, ws.max_column + 1)]
            vals = [c for c in cells if c.value is not None]
            if not vals:
                continue
            lab = label_of(ws, r)
            forms = [c for c in vals if is_formula(c.value)]
            nums = [c for c in vals if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
            pats = Counter(X.relative_form(c.value, c.coordinate) for c in forms)
            desc = []
            if forms:
                first = forms[0]
                desc.append(f"{len(forms)}f {first.coordinate}:{first.value[:70]}")
                if len(pats) > 1:
                    desc.append(f"{len(pats)} patterns")
            if nums:
                blue = sum(1 for c in nums if font_rgb(c) == "0000FF")
                desc.append(f"{len(nums)}n" + (f" ({blue} blue)" if blue else "")
                            + f" {nums[0].coordinate}={nums[0].value!r}")
            sample = []
            if forms:
                for c in forms[-1:]:
                    sample.append(f"{c.coordinate}->{wsv[c.coordinate].value!r}")
            print(f"r{r:<4} {lab[:38]:<38} | " + " | ".join(desc) + (f" | {', '.join(sample)}" if sample else ""))
            shown += 1
            if shown >= a.max_rows:
                print(f"... truncated at {a.max_rows} rows (use --max-rows)")
                break


# ============================================================================ audit

def cmd_audit(a):
    import audit
    audit.main(a)


def _tokens(f):
    try:
        return X.Tokenizer(f).items
    except Exception:
        return []


ARG_EXEMPT = {"ROUND", "ROUNDUP", "ROUNDDOWN", "MROUND", "INDEX", "MATCH", "VLOOKUP", "HLOOKUP", "CHOOSE",
              "OFFSET", "LEFT", "RIGHT", "MID", "TEXT", "DATE", "EDATE", "EOMONTH", "LARGE", "SMALL", "PERCENTILE",
              "QUARTILE", "RANK", "WEEKDAY", "XNPV", "XIRR", "IRR", "_XLFN.XLOOKUP", "_XLFN.XMATCH", "SUBTOTAL", "AGGREGATE"}


def _literal_numbers(f):
    """Numbers typed into a formula, ignoring arguments that are structural rather than
    economic (ROUND digits, INDEX/MATCH positions, DATE parts and so on)."""
    stack, out = [], []
    for t in _tokens(f):
        if t.type == "FUNC" and t.subtype == "OPEN":
            stack.append(t.value[:-1].upper())
        elif t.type == "FUNC" and t.subtype == "CLOSE":
            if stack:
                stack.pop()
        elif t.type == "OPERAND" and t.subtype == "NUMBER":
            if stack and stack[-1] in ARG_EXEMPT:
                continue
            try:
                out.append(float(t.value))
            except ValueError:
                pass
    return out


def _only_link(f):
    toks = [t for t in _tokens(f) if t.type != "WHITE-SPACE"]
    return len(toks) == 1 and toks[0].subtype == "RANGE"


def expand(addr, maxcells=5000):
    try:
        c1, r1, c2, r2 = range_boundaries(addr)
    except Exception:
        return []
    if None in (c1, r1, c2, r2) or (c2 - c1 + 1) * (r2 - r1 + 1) > maxcells:
        return []
    return [f"{i2col(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def graph(wf):
    names = {}
    for n, dn in wf.defined_names.items():
        names[n] = dn.attr_text
    g = {}
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not is_formula(c.value):
                    continue
                deps = set()
                for sh, addr in X.iter_formula_refs(c.value):
                    if addr in names and not sh:
                        sh, addr = X.split_sheet(names[addr])
                    for cell in expand(addr.replace("$", "")):
                        deps.add(f"{sh or ws.title}!{cell}")
                g[f"{ws.title}!{c.coordinate}"] = deps
    return g


def find_cycles(wf):
    g = graph(wf)
    color, cycles = {}, []
    sys.setrecursionlimit(100000)

    def dfs(u, stack):
        color[u] = 1
        stack.append(u)
        for v in g.get(u, ()):
            if color.get(v) == 1:
                cycles.append(" -> ".join(stack[stack.index(v):] + [v])[:300])
            elif color.get(v) is None and v in g:
                dfs(v, stack)
        stack.pop()
        color[u] = 2

    for u in list(g):
        if color.get(u) is None:
            dfs(u, [])
    return cycles


# ============================================================================ patch engine

class Patcher:
    def __init__(self, path):
        self.pkg = X.Package(path)
        self.touched = []          # (sheet, cell) written
        self.formula_changed = False
        self._sst = None

    # -- reading values from XML (for label/header lookups)
    def sst(self):
        if self._sst is None:
            self._sst = []
            if "xl/sharedStrings.xml" in self.pkg.raw:
                for si in self.pkg.xml("xl/sharedStrings.xml").findall("m:si", NS):
                    self._sst.append("".join(si.itertext()))
        return self._sst

    def cell_text(self, c):
        t = c.get("t")
        if t == "s":
            v = c.find("m:v", NS)
            return self.sst()[int(v.text)] if v is not None else None
        if t == "inlineStr":
            return "".join(c.find("m:is", NS).itertext())
        v = c.find("m:v", NS)
        return v.text if v is not None else None

    def resolve(self, sheet, spec):
        """'B5' or {'row': 'Revenue', 'col': 'FY2027'} -> 'B5'. Labels are matched in
        columns A-C, headers in rows 1-15, case- and whitespace-insensitive."""
        if isinstance(spec, str):
            return spec.replace("$", "")
        if "key" in spec:
            # rows built by model_builder carry a defined name k_<Sheet>_<key> on their label
            # cell; Excel moves the name when rows are inserted, so this survives user edits
            want = f"k_{re.sub(r'[^A-Za-z0-9_]', '_', sheet)}_{spec['key']}"
            dn = self.pkg.xml("xl/workbook.xml").find("m:definedNames", NS)
            for d in (dn if dn is not None else []):
                if d.get("name") == want:
                    _, addr = X.split_sheet(d.text)
                    _, r = X.split_cell(addr.replace("$", ""))
                    col = spec.get("col", "C")
                    if not re.fullmatch(r"[A-Z]{1,3}", col):
                        col = re.sub(r"\d+", "", self.resolve(sheet, {"row": r, "col": col}))
                    return f"{col}{r}"
            raise KeyError(f"no defined name {want}")
        root = self.pkg.xml(self.pkg.sheet_part(sheet))
        norm = lambda s: re.sub(r"\s+", " ", str(s or "")).strip().lower()
        row = spec.get("row")
        col = spec.get("col")
        rnum = row if isinstance(row, int) else None
        cnum = col2i(col) if isinstance(col, str) and re.fullmatch(r"[A-Z]{1,3}", col) else None
        for c in X.iter_cells(root):
            ci, ri = X.split_cell(c.get("r"))
            txt = norm(self.cell_text(c))
            if rnum is None and ci <= 3 and txt == norm(row):
                rnum = ri
            if cnum is None and ri <= 15 and txt in (norm(col), norm(col).rstrip(".0")):
                cnum = ci
        if rnum is None or cnum is None:
            raise KeyError(f"could not resolve {spec} on {sheet} (row={rnum}, col={cnum})")
        return f"{i2col(cnum)}{rnum}"

    def _root(self, sheet):
        part = self.pkg.sheet_part(sheet)
        self.pkg.touch(part)
        return self.pkg.xml(part)

    def _guard_shared(self, root, cell):
        f = cell.find("m:f", NS)
        if f is not None and f.get("t") == "shared" and f.get("ref"):
            X.expand_shared_formulas(root, only_si=f.get("si"))
        if f is not None and f.get("t") == "array" and f.get("ref") and ":" in f.get("ref"):
            raise ValueError(f"{cell.get('r')} is the anchor of a multi-cell array formula {f.get('ref')}; edit the array as a whole")
        for c in X.iter_cells(root):
            ff = c.find("m:f", NS)
            if ff is not None and ff.get("t") == "array" and ff.get("ref") and ":" in ff.get("ref") and c is not cell:
                if cell.get("r") in expand(ff.get("ref")):
                    raise ValueError(f"{cell.get('r')} is inside array formula {ff.get('ref')} anchored at {c.get('r')}")

    # -- operations
    def op_set(self, o):
        sheet = o["sheet"]
        root = self._root(sheet)
        ref = self.resolve(sheet, o["cell"])
        cell = X.find_cell(root, ref, create=True)
        self._guard_shared(root, cell)
        if "style_like" in o:
            like = X.find_cell(root, self.resolve(sheet, o["style_like"]))
            if like is not None and like.get("s"):
                cell.set("s", like.get("s"))
        if "formula" in o:
            X.set_cell(cell, formula=o["formula"])
            self.formula_changed = True
        else:
            had_f = cell.find("m:f", NS) is not None
            X.set_cell(cell, value=o.get("value"))
            self.formula_changed |= had_f
        self.touched.append((sheet, ref))

    def op_fill(self, o):
        """Copy the formula in `src` across `range`, adjusting relative refs (Ctrl+R)."""
        sheet = o["sheet"]
        root = self._root(sheet)
        src = self.resolve(sheet, o["src"])
        sc = X.find_cell(root, src)
        f = sc.find("m:f", NS) if sc is not None else None
        if f is None:
            raise ValueError(f"{sheet}!{src} has no formula to fill")
        if f.get("t") == "shared" and not f.text:
            X.expand_shared_formulas(root, only_si=f.get("si"))
            f = sc.find("m:f", NS)
        text = "=" + f.text
        for ref in expand(o["range"].replace("$", "")):
            if ref == src:
                continue
            cell = X.find_cell(root, ref, create=True)
            self._guard_shared(root, cell)
            if sc.get("s") and o.get("copy_style", True):
                cell.set("s", sc.get("s"))
            X.set_cell(cell, formula=Translator(text, origin=src).translate_formula(ref))
            self.touched.append((sheet, ref))
        self.formula_changed = True

    def op_style(self, o):
        sheet = o["sheet"]
        root = self._root(sheet)
        like = X.find_cell(root, self.resolve(sheet, o["like"]))
        s = like.get("s") if like is not None else None
        for ref in expand(o["range"].replace("$", "")):
            cell = X.find_cell(root, ref, create=True)
            if s:
                cell.set("s", s)
            else:
                cell.attrib.pop("s", None)

    def op_name(self, o):
        wb = self.pkg.xml("xl/workbook.xml")
        dn = wb.find("m:definedNames", NS)
        if dn is None:
            dn = etree.Element(M + "definedNames")
            wb.find("m:sheets", NS).addnext(dn)
        for d in dn.findall("m:definedName", NS):
            if d.get("name") == o["name"] and d.get("localSheetId") is None:
                dn.remove(d)
        d = etree.SubElement(dn, M + "definedName")
        d.set("name", o["name"])
        d.text = o["ref"].lstrip("=")
        self.pkg.touch("xl/workbook.xml")
        self.formula_changed = True

    def op_iterate(self, o):
        self.pkg.force_full_calc(iterate=bool(o.get("on", True)))

    def op_insert(self, o, axis):
        insert_lines(self.pkg, o["sheet"], axis, o["at"], int(o.get("n", 1)), o.get("copy_from"))
        self.formula_changed = True

    def apply(self, ops):
        for o in ops:
            kind = o["op"]
            if kind == "set":
                self.op_set(o)
            elif kind == "fill":
                self.op_fill(o)
            elif kind == "style":
                self.op_style(o)
            elif kind == "name":
                self.op_name(o)
            elif kind == "iterate":
                self.op_iterate(o)
            elif kind == "insert_rows":
                self.op_insert(o, "row")
            elif kind == "insert_cols":
                self.op_insert(o, "col")
            else:
                raise ValueError(f"unknown op {kind}")
        if self.formula_changed or self.touched:
            self.pkg.drop_calc_chain()
            self.pkg.force_full_calc()


def insert_lines(pkg, sheet, axis, at, n, copy_from=None):
    """Insert n rows (axis='row', at=row number) or columns (axis='col', at=letter or
    number) before `at`, rewriting every reference in the workbook the way Excel does."""
    if axis == "col" and isinstance(at, str):
        at = col2i(at)
    if axis == "col" and isinstance(copy_from, str):
        copy_from = col2i(copy_from)
    target = pkg.sheet_part(sheet)
    sheets = pkg.sheets()
    # 1. explicit formulas everywhere (shared formula groups cannot be shifted safely)
    for name, part, _ in sheets:
        root = pkg.xml(part)
        if X.expand_shared_formulas(root):
            pkg.touch(part)

    def shift_coord(ref):
        c, r = X.split_cell(ref)
        if axis == "row" and r >= at:
            r += n
        if axis == "col" and c >= at:
            c += n
        return f"{i2col(c)}{r}"

    # 2. every formula in every sheet
    for name, part, _ in sheets:
        root = pkg.xml(part)
        for f in root.iter(M + "f"):
            if f.text:
                new = X.shift_formula(f.text, name, sheet, axis, at, n)
                if new != f.text:
                    f.text = new
                    pkg.touch(part)
        for el in root.iter("{%s}f" % NS["xm"]):
            if el.text:
                new = X.shift_formula(el.text, name, sheet, axis, at, n)
                if new != el.text:
                    el.text = new
                    pkg.touch(part)
    # 3. the target sheet's own geometry
    root = pkg.xml(target)
    pkg.touch(target)
    sd = X.sheet_data(root)
    rows = sd.findall("m:row", NS)
    for row in reversed(rows):
        r = int(row.get("r"))
        if axis == "row" and r >= at:
            row.set("r", str(r + n))
        for c in row.findall("m:c", NS):
            c.set("r", shift_coord(c.get("r")))
        if axis == "col":
            row.attrib.pop("spans", None)
    for f in root.iter(M + "f"):
        for attr in ("ref", "r1", "r2"):
            if f.get(attr):
                f.set(attr, X.shift_sqref(f.get(attr), axis, at, n))
    for tag in ("conditionalFormatting", "dataValidation"):
        for el in root.iter(M + tag):
            if el.get("sqref"):
                el.set("sqref", X.shift_sqref(el.get("sqref"), axis, at, n, extend_adjacent=True))
    for tag, attr in (("dimension", "ref"), ("mergeCell", "ref"), ("hyperlink", "ref"), ("autoFilter", "ref"),
                      ("selection", "sqref"), ("selection", "activeCell"), ("pane", "topLeftCell"),
                      ("protectedRange", "sqref")):
        for el in root.iter(M + tag):
            if el.get(attr):
                el.set(attr, X.shift_sqref(el.get(attr), axis, at, n))
    for tag in ("formula", "formula1", "formula2"):
        for el in root.iter(M + tag):
            if el.text:
                el.text = X.shift_formula(el.text, sheet, sheet, axis, at, n)
    for el in root.iter("{%s}sqref" % NS["xm"]):
        if el.text:
            el.text = X.shift_sqref(el.text, axis, at, n)
    if axis == "row":
        for brk in root.iter(M + "brk"):
            if brk.getparent().tag == M + "rowBreaks" and int(brk.get("id")) >= at:
                brk.set("id", str(int(brk.get("id")) + n))
    if axis == "col":
        cols = root.find("m:cols", NS)
        if cols is not None:
            for col in cols.findall("m:col", NS):
                mn, mx = int(col.get("min")), int(col.get("max"))
                if mn >= at:
                    col.set("min", str(mn + n)); col.set("max", str(mx + n))
                elif mx >= at:
                    col.set("max", str(mx + n))
    # 4. copy a template row/col into the gap (style + translated formulas)
    if copy_from is not None:
        src_idx = copy_from + (n if copy_from >= at else 0)
        if axis == "row":
            src_row = X.find_row(root, src_idx)
            if src_row is not None:
                for k in range(n):
                    rr = at + k
                    new_row = X.find_row(root, rr, create=True)
                    for a_ in ("s", "customFormat", "ht", "customHeight", "outlineLevel"):
                        if src_row.get(a_):
                            new_row.set(a_, src_row.get(a_))
                    for c in src_row.findall("m:c", NS):
                        ci, _ = X.split_cell(c.get("r"))
                        ref = f"{i2col(ci)}{rr}"
                        nc = X.find_cell(root, ref, create=True)
                        if c.get("s"):
                            nc.set("s", c.get("s"))
                        f = c.find("m:f", NS)
                        if f is not None and f.text and not f.get("t"):
                            X.set_cell(nc, formula=Translator("=" + f.text, origin=c.get("r")).translate_formula(ref))
        else:
            for row in sd.findall("m:row", NS):
                src = X.find_cell(root, f"{i2col(src_idx)}{row.get('r')}")
                if src is None:
                    continue
                for k in range(n):
                    ref = f"{i2col(at + k)}{row.get('r')}"
                    nc = X.find_cell(root, ref, create=True)
                    if src.get("s"):
                        nc.set("s", src.get("s"))
                    f = src.find("m:f", NS)
                    if f is not None and f.text and not f.get("t"):
                        X.set_cell(nc, formula=Translator("=" + f.text, origin=src.get("r")).translate_formula(ref))
    # 5. parts hanging off the target sheet: tables, drawings, comments
    for rid, (typ, tpart) in pkg.rels(target).items():
        if tpart not in pkg.raw:
            continue
        if typ == "table":
            t = pkg.xml(tpart)
            t.set("ref", X.shift_sqref(t.get("ref"), axis, at, n))
            for af in t.iter(M + "autoFilter"):
                af.set("ref", X.shift_sqref(af.get("ref"), axis, at, n))
            pkg.touch(tpart)
        elif typ == "drawing":
            d = pkg.xml(tpart)
            tag = "row" if axis == "row" else "col"
            for anchor in d:
                for side in ("from", "to"):
                    s = anchor.find(f"xdr:{side}", NS)
                    if s is not None:
                        el = s.find(f"xdr:{tag}", NS)
                        if el is not None and int(el.text) + 1 >= at:
                            el.text = str(int(el.text) + n)
            pkg.touch(tpart)
        elif typ == "comments":
            cm = pkg.xml(tpart)
            for c in cm.iter(M + "comment"):
                c.set("ref", shift_coord(c.get("ref")))
            pkg.touch(tpart)
        elif typ == "vmlDrawing":
            txt = pkg.raw[tpart].decode("utf8", "ignore")
            def fix_row(m):
                v = int(m.group(2))
                return f"{m.group(1)}{v + n if v + 1 >= at else v}{m.group(3)}"
            if axis == "row":
                txt = re.sub(r"(<x:Row>)(\d+)(</x:Row>)", fix_row, txt)
            else:
                txt = re.sub(r"(<x:Column>)(\d+)(</x:Column>)", fix_row, txt)
            pkg.raw[tpart] = txt.encode("utf8")
    # 6. workbook-level references: defined names, charts, pivot caches
    wb = pkg.xml("xl/workbook.xml")
    dn = wb.find("m:definedNames", NS)
    if dn is not None:
        for d in dn:
            if d.text:
                new = X.shift_formula(d.text, "\x00", sheet, axis, at, n)
                if new != d.text:
                    d.text = new
                    pkg.touch("xl/workbook.xml")
    for part in pkg.names():
        if part.startswith("xl/charts/chart") and part.endswith(".xml"):
            ch = pkg.xml(part)
            for el in ch.iter("{%s}f" % NS["c"]):
                if el.text:
                    new = X.shift_formula(el.text, "\x00", sheet, axis, at, n)
                    if new != el.text:
                        el.text = new
                        pkg.touch(part)
        if part.startswith("xl/pivotCache/pivotCacheDefinition"):
            pc = pkg.xml(part)
            for ws_ in pc.iter(M + "worksheetSource"):
                if ws_.get("sheet") == sheet and ws_.get("ref"):
                    ws_.set("ref", X.shift_sqref(ws_.get("ref"), axis, at, n))
                    pkg.touch(part)


def refresh_cache(path, recalced):
    """Write LibreOffice-computed values into the <v> of every formula cell in `path`, so
    viewers that never calculate (previewers, pandas, GitHub) show current numbers. Excel
    still recalculates on open (fullCalcOnLoad)."""
    pkg = X.Package(path)
    wv = load_workbook(recalced, data_only=True)
    n = 0
    for name, part, _ in pkg.sheets():
        if name not in wv.sheetnames:
            continue
        ws = wv[name]
        root = pkg.xml(part)
        for c in X.iter_cells(root):
            f = c.find("m:f", NS)
            if f is None or c.get("cm"):
                continue
            val = ws[c.get("r")].value
            for v in c.findall("m:v", NS):
                c.remove(v)
            c.attrib.pop("t", None)
            if val is None:
                continue
            v = etree.SubElement(c, M + "v")
            if isinstance(val, bool):
                c.set("t", "b"); v.text = "1" if val else "0"
            elif isinstance(val, (int, float)):
                v.text = repr(float(val)) if isinstance(val, float) else str(val)
            elif isinstance(val, str) and val in ERRORS:
                c.set("t", "e"); v.text = val
            elif hasattr(val, "isoformat"):
                from openpyxl.utils.datetime import to_excel
                v.text = repr(float(to_excel(val)))
            else:
                c.set("t", "str"); v.text = str(val)
            n += 1
        pkg.touch(part)
    pkg.save()
    return n


def verify(path, cells=None):
    rc = X.lo_recalc_copy(path)
    wv = load_workbook(rc, data_only=True)
    wf = load_workbook(path)
    wl = load_workbook(rc)          # LibreOffice's own formula text: lowercased = it could not parse it
    unparsed = []
    errs = defaultdict(list)
    total = 0
    for ws in wf.worksheets:
        wsv = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if is_formula(c.value) or type(c.value).__name__ == "ArrayFormula":
                    total += 1
                    v = wsv[c.coordinate].value
                    if isinstance(v, str) and v in ERRORS:
                        errs[v].append(f"{ws.title}!{c.coordinate}")
                    lf = wl[ws.title][c.coordinate].value
                    if isinstance(lf, str) and re.search(r"\b[a-z][a-z0-9.]*\(", lf) and \
                            not re.search(r"\b[a-z][a-z0-9.]*\(", str(c.value)):
                        unparsed.append(f"{ws.title}!{c.coordinate}: {lf[:60]}")
    vals = {}
    for spec in cells or []:
        sh, addr = parse_cellspec(spec)
        vals[spec] = wv[sh][addr].value
    ext = [n for n in X.Package(path).names() if n.startswith("xl/externalLinks/")]
    return {"recalculated_copy": str(rc), "formulas": total, "external_link_parts": len(ext),
            "libreoffice_could_not_parse": unparsed[:25],
            "errors": {k: {"count": len(v), "cells": v[:25]} for k, v in errs.items()},
            "status": "errors_found" if errs else ("unverifiable_formulas" if unparsed else "clean"), "values": vals}


def cmd_patch(a):
    ops = json.loads(Path(a.ops).read_text()) if a.ops.endswith(".json") else json.loads(a.ops)
    p = Patcher(a.file)
    p.apply(ops)
    out = p.pkg.save(a.out)
    rep = {"out": str(out), "cells_written": len(p.touched)}
    if not a.no_verify:
        v = verify(out, [f"{X.quote_sheet(s)}!{c}" for s, c in p.touched[:40]])
        rep["verify"] = v
        if a.refresh_cache:
            if v["external_link_parts"] and not a.force:
                # LibreOffice cannot open the linked files: it would cache #NAME?/#REF! over the only copy of that data
                rep["cached_values_written"] = "REFUSED: workbook links to other files; cached values are the only copy (use --force to override)"
            else:
                rep["cached_values_written"] = refresh_cache(out, v["recalculated_copy"])
    print(json.dumps(rep, indent=1, default=str))


def cmd_verify(a):
    print(json.dumps(verify(a.file, [c.strip() for c in a.cells.split(",")] if a.cells else []), indent=1, default=str))


# ============================================================================ diff / trace

def cmd_diff(a):
    fa, fb = load_workbook(a.old), load_workbook(a.new)
    va, vb = load_workbook(a.old, data_only=True), load_workbook(a.new, data_only=True)
    expect = set()
    for spec in (a.expect.split(",") if a.expect else []):
        sh, addr = parse_cellspec(spec.strip())
        for c in expand(addr) or [addr]:
            expect.add(f"{sh}!{c}")
    changes, unexpected = [], []
    for name in sorted(set(fa.sheetnames) | set(fb.sheetnames)):
        if name not in fa.sheetnames or name not in fb.sheetnames:
            changes.append(f"sheet {name} {'added' if name in fb.sheetnames else 'removed'}")
            continue
        A, B = fa[name], fb[name]
        mr, mc = max(A.max_row, B.max_row), max(A.max_column, B.max_column)
        for r in range(1, mr + 1):
            for c in range(1, mc + 1):
                x, y = A.cell(r, c).value, B.cell(r, c).value
                vx, vy = va[name].cell(r, c).value, vb[name].cell(r, c).value
                key = f"{name}!{i2col(c)}{r}"
                fchg = str(x) != str(y)
                vchg = not _close(vx, vy)
                if fchg or (vchg and a.values):
                    line = f"{key}: {x!r} -> {y!r}" if fchg else f"{key}: value {vx!r} -> {vy!r}"
                    changes.append(line)
                    if expect and fchg and key not in expect:
                        unexpected.append(line)
    print(json.dumps({"changed": len(changes), "changes": changes[: a.limit],
                      "unexpected_input_or_formula_changes": unexpected[: a.limit]}, indent=1, default=str))


def _close(x, y):
    if isinstance(x, (int, float)) and isinstance(y, (int, float)):
        return abs(x - y) <= 1e-9 * max(1, abs(x), abs(y))
    return x == y


def cmd_trace(a):
    wf = load_workbook(a.file)
    wv = load_workbook(X.lo_recalc_copy(a.file), data_only=True)
    names = {n: dn.attr_text for n, dn in wf.defined_names.items()}
    sh, addr = parse_cellspec(a.cell)
    if a.dependents:
        g = graph(wf)
        rev = defaultdict(set)
        for k, deps in g.items():
            for d in deps:
                rev[d].add(k)
        seen, frontier = set(), {f"{sh}!{addr}"}
        for depth in range(a.depth):
            nxt = set()
            for u in frontier:
                for v in rev.get(u, ()):
                    if v not in seen:
                        seen.add(v); nxt.add(v)
                        s2, a2 = v.split("!")
                        print(f"{'  ' * depth}<- {v} = {wf[s2][a2].value} [{wv[s2][a2].value!r}] ({label_of(wf[s2], wf[s2][a2].row)})")
            frontier = nxt
        return

    def walk(sheet, ref, depth, seen):
        key = f"{sheet}!{ref}"
        c, cv = wf[sheet][ref], wv[sheet][ref]
        print(f"{'  ' * depth}{key} ({label_of(wf[sheet], c.row)}) = {c.value!r} [{cv.value!r}]")
        if depth >= a.depth or key in seen or not is_formula(c.value):
            return
        seen.add(key)
        for s2, ad in X.iter_formula_refs(c.value):
            if ad in names and not s2:
                s2, ad = X.split_sheet(names[ad])
            cells = expand(ad.replace("$", ""), 50)
            if len(cells) > 6:
                print(f"{'  ' * (depth + 1)}{s2 or sheet}!{ad} (range of {len(cells)})")
                continue
            for cc in cells:
                walk(s2 or sheet, cc, depth + 1, seen)

    walk(sh, addr, 0, set())


def sensitivity(path, out_cell, row_input, row_vals, col_input=None, col_vals=None):
    """Brute-force a 1D/2D sensitivity: set the input cell(s), recalculate the whole model,
    read the output. Unlike an Excel Data Table this is verifiable here, and unlike an
    algebraic shortcut it respects every non-linearity in the model."""
    grid = []
    with X.LOSession(path) as lo:
        base_r = lo.get(row_input)
        base_c = lo.get(col_input) if col_input else None
        for cv in (col_vals or [None]):
            line = []
            if col_input:
                lo.set(col_input, cv)
            for rv in row_vals:
                lo.set(row_input, rv)
                lo.recalc()
                line.append(lo.get(out_cell))
            grid.append(line)
        lo.set(row_input, base_r)
        if col_input:
            lo.set(col_input, base_c)
        lo.recalc()
        base_out = lo.get(out_cell)
    return {"output": out_cell, "row_input": row_input, "row_values": row_vals,
            "col_input": col_input, "col_values": col_vals, "grid": grid, "base_output": base_out}


def grid_assertions(res):
    """Sell-side conventions for a sensitivity grid: odd size with a true centre, the
    centre equals the live model, every cell distinct, values monotonic along each axis."""
    g = res["grid"]
    out = {}
    nums = [v for row in g for v in row if isinstance(v, (int, float))]
    out["all_numeric"] = len(nums) == sum(len(r) for r in g)
    out["odd_dimensions"] = len(res["row_values"]) % 2 == 1 and len(res["col_values"] or [0]) % 2 == 1
    ri = len(res["row_values"]) // 2
    ci = len(res["col_values"] or [0]) // 2
    centre = g[ci][ri] if out["all_numeric"] else None
    out["centre_equals_base"] = centre is not None and _close(centre, res["base_output"])
    out["all_distinct"] = len({round(v, 9) for v in nums}) == len(nums)

    def mono(seq):
        d = [b - a for a, b in zip(seq, seq[1:])]
        return "up" if all(x > 0 for x in d) else "down" if all(x < 0 for x in d) else "mixed"
    if out["all_numeric"]:
        out["direction_along_row_input"] = sorted({mono(row) for row in g})
        if len(g) > 1:
            out["direction_along_col_input"] = sorted({mono([row[j] for row in g]) for j in range(len(g[0]))})
    out["pass"] = all(v for k, v in out.items() if isinstance(v, bool)) and \
        "mixed" not in out.get("direction_along_row_input", []) + out.get("direction_along_col_input", [])
    return out


def scenarios(path, switch, cases, outputs, check=None):
    rows = []
    with X.LOSession(path) as lo:
        base = lo.get(switch)
        for case in cases:
            lo.set(switch, case)
            lo.recalc()
            row = {"case": case, **{o: lo.get(o) for o in outputs}}
            if check:
                row["check"] = lo.get(check)
            rows.append(row)
        lo.set(switch, base)
        lo.recalc()
    return rows


def cmd_scenarios(a):
    sw, cases = a.switch.split("=", 1)
    cases = cases.split(",")
    outs = [o.strip() for o in a.outputs.split(",")]
    rows = scenarios(a.file, sw, cases, outs, a.check)
    res = {"switch": sw, "rows": rows, "assertions": {}}
    # a switch that moves nothing is the classic silent failure (wrong cell, or a hard-coded live column)
    res["assertions"]["switch_changes_outputs"] = any(
        not _close(rows[0][o], r[o]) for r in rows[1:] for o in outs if isinstance(r[o], (int, float)))
    if a.check:
        res["assertions"]["master_check_passes_in_every_case"] = all(r["check"] == 1 for r in rows)
    for spec in (a.order or "").split(","):
        if not spec.strip():
            continue
        cell, direction = spec.strip().rsplit(":", 1)
        vals = [r[cell] for r in rows]
        ok = all((b > a_) if direction == "asc" else (b < a_) for a_, b in zip(vals, vals[1:]))
        res["assertions"][f"{cell} {direction} across {','.join(cases)}"] = ok
    print(json.dumps(res, indent=1, default=str))
    if not all(res["assertions"].values()):
        sys.exit(3)


def tornado(path, output, specs):
    """One-at-a-time flex of each input; ranked by swing |high - low|."""
    res = []
    with X.LOSession(path) as lo:
        base_out = lo.get(output)
        for cell, lo_v, hi_v in specs:
            b = lo.get(cell)
            lo.set(cell, lo_v); lo.recalc(); v_lo = lo.get(output)
            lo.set(cell, hi_v); lo.recalc(); v_hi = lo.get(output)
            lo.set(cell, b); lo.recalc()
            res.append({"input": cell, "base_input": b, "low": lo_v, "high": hi_v,
                        "output_at_low": v_lo, "output_at_high": v_hi,
                        "swing": abs(v_hi - v_lo) if all(isinstance(v, (int, float)) for v in (v_hi, v_lo)) else None})
    res.sort(key=lambda r: -(r["swing"] or 0))
    return {"output": output, "base_output": base_out, "ranked": res}


def cmd_tornado(a):
    wf = load_workbook(a.file)
    specs = []
    for item in a.inputs.split(";"):
        item = item.strip()
        if "=" in item:
            cell, rng = item.split("=", 1)
            lo_v, hi_v = (float(x) for x in rng.split(","))
        else:
            cell = item
            sh, addr = parse_cellspec(cell)
            b = wf[sh][addr].value
            if not isinstance(b, (int, float)):
                raise SystemExit(f"{cell} is not a numeric input; give an explicit range like {cell}=lo,hi")
            lo_v, hi_v = b * (1 - a.pct), b * (1 + a.pct)
        specs.append((cell, lo_v, hi_v))
    res = tornado(a.file, a.output, specs)
    for r in res["ranked"]:
        sh, addr = parse_cellspec(r["input"])
        r["label"] = label_of(wf[sh], wf[sh][addr].row)
    print(json.dumps(res, indent=1, default=str))


def _parse_axis(s):
    cell, vals = s.split("=", 1)
    return cell, [float(v) for v in vals.split(",")]


def cmd_sensitivity(a):
    ri, rv = _parse_axis(a.row)
    ci, cv = _parse_axis(a.col) if a.col else (None, None)
    res = sensitivity(a.file, a.output, ri, rv, ci, cv)
    res["assertions"] = grid_assertions(res)
    if a.write:
        # write the verified grid as static values with a top-left anchor, plus a provenance note
        sh, anchor = parse_cellspec(a.write)
        c0, r0 = X.split_cell(anchor)
        ops = []
        for j, v in enumerate(rv):
            ops.append({"op": "set", "sheet": sh, "cell": f"{i2col(c0 + 1 + j)}{r0}", "value": v})
        for i, v in enumerate(cv or [None]):
            if v is not None:
                ops.append({"op": "set", "sheet": sh, "cell": f"{i2col(c0)}{r0 + 1 + i}", "value": v})
            for j, out in enumerate(res["grid"][i]):
                ops.append({"op": "set", "sheet": sh, "cell": f"{i2col(c0 + 1 + j)}{r0 + 1 + i}", "value": out})
        ops.append({"op": "set", "sheet": sh, "cell": f"{i2col(c0)}{r0}",
                    "value": f"{a.output} by {ri}" + (f" x {ci}" if ci else "") + " (static; rerun xl.py sensitivity)"})
        p = Patcher(a.file); p.apply(ops); p.pkg.save(a.out or a.file)
        res["written_to"] = a.out or a.file
    print(json.dumps(res, indent=1, default=str))


def cmd_render(a):
    """Render every sheet (charts included) of a recalculated copy to PNG pages, so the
    layout and charts can be looked at. Uses LibreOffice PDF export plus pdftoppm."""
    import subprocess, shutil
    out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    rc = X.lo_recalc_copy(a.file)
    prof = Path(rc).parent.parent / "profile"
    subprocess.run(["soffice", f"-env:UserInstallation={prof.as_uri()}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(out), str(rc)], capture_output=True, timeout=180)
    pdf = out / (Path(rc).stem + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-r", str(a.dpi), str(pdf), str(out / "page")], check=True)
    pages = sorted(str(p) for p in out.glob("page*.png"))
    print(json.dumps({"pdf": str(pdf), "pages": pages}, indent=1))


def cmd_features(a):
    print(json.dumps(features(a.file), indent=1))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("map"); p.add_argument("file"); p.add_argument("--sheet"); p.add_argument("--max-rows", type=int, default=120)
    p.add_argument("--no-recalc", action="store_true", help="show cached values instead of a fresh LibreOffice recalc")
    p = sp.add_parser("audit"); p.add_argument("file"); p.add_argument("--sheet")
    p.add_argument("--format", choices=["json", "md"], default="json"); p.add_argument("--limit", type=int, default=300)
    p.add_argument("--annotate", help="write a REVIEW copy with comments and fills on every finding")
    p.add_argument("--suppress", help='JSON list of accepted findings [{"id":..,"sheet":..,"cell":..}]')
    p.add_argument("--fail-on", choices=["high", "medium", "low"], help="exit 2 if any finding at or above this severity")
    p = sp.add_parser("patch"); p.add_argument("file"); p.add_argument("ops"); p.add_argument("--out", required=True)
    p.add_argument("--refresh-cache", action="store_true"); p.add_argument("--no-verify", action="store_true")
    p.add_argument("--force", action="store_true", help="refresh cached values even when external links exist")
    p = sp.add_parser("verify"); p.add_argument("file"); p.add_argument("--cells")
    p = sp.add_parser("diff"); p.add_argument("old"); p.add_argument("new"); p.add_argument("--expect")
    p.add_argument("--values", action="store_true", help="also list cached-value changes"); p.add_argument("--limit", type=int, default=200)
    p = sp.add_parser("trace"); p.add_argument("file"); p.add_argument("cell"); p.add_argument("--depth", type=int, default=3)
    p.add_argument("--dependents", action="store_true")
    p = sp.add_parser("features"); p.add_argument("file")
    p = sp.add_parser("sensitivity"); p.add_argument("file"); p.add_argument("output", help="e.g. Model!F9")
    p.add_argument("--row", required=True, help="Inputs!B3=0,0.05,0.1"); p.add_argument("--col", help="Inputs!B4=0.5,0.55")
    p.add_argument("--write", help="anchor cell for a static grid, e.g. Outputs!B20"); p.add_argument("--out")
    p = sp.add_parser("scenarios"); p.add_argument("file"); p.add_argument("--switch", required=True, help="Inputs!C4=Bear,Base,Bull")
    p.add_argument("--outputs", required=True); p.add_argument("--check", help="master check cell, must be 1")
    p.add_argument("--order", help="Valuation!C10:asc  (strictly increasing in case order)")
    p = sp.add_parser("render"); p.add_argument("file"); p.add_argument("--out", required=True); p.add_argument("--dpi", type=int, default=70)
    p = sp.add_parser("tornado"); p.add_argument("file"); p.add_argument("output")
    p.add_argument("--inputs", required=True, help='"Inputs!C17=0.08,0.12; Inputs!C14" (bare cell = base +/- pct)')
    p.add_argument("--pct", type=float, default=0.10)
    a = ap.parse_args()
    {"map": cmd_map, "audit": cmd_audit, "patch": cmd_patch, "verify": cmd_verify, "diff": cmd_diff,
     "trace": cmd_trace, "features": cmd_features, "sensitivity": cmd_sensitivity,
     "scenarios": cmd_scenarios, "tornado": cmd_tornado, "render": cmd_render}[a.cmd](a)


if __name__ == "__main__":
    main()
