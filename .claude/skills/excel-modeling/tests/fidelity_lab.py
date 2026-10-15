"""Reproduce the round-trip fidelity table in references/fidelity.md.

Builds a workbook with the features financial models use, then saves it three ways:
openpyxl load+save, LibreOffice re-save, and xl.py patch (XML-level). Prints which
features survive each path. Run: python3 tests/fidelity_lab.py
"""
import os, re, sys, tempfile, warnings, zipfile
warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "scripts"))
import openpyxl, xlsxwriter
import xl, xlcore as X

W = tempfile.mkdtemp(prefix="fidelity_")
SRC = os.path.join(W, "sink.xlsx")

wb = xlsxwriter.Workbook(SRC)
ws = wb.add_worksheet("Inputs"); calc = wb.add_worksheet("Model Calc")
blue = wb.add_format({"font_color": "blue", "num_format": "$#,##0;($#,##0);-"})
ws.write("A1", "Growth"); ws.write("B1", 0.05); ws.write("A2", "Base"); ws.write("B2", 1000, blue)
wb.define_name("Growth", "=Inputs!$B$1")
ws.data_validation("B1", {"validate": "decimal", "criteria": "between", "minimum": -1, "maximum": 1})
ws.conditional_format("B1", {"type": "cell", "criteria": ">", "value": 0.1, "format": blue})
ws.write_comment("B2", "Source note")
calc.write_row("A1", ["Year", 2025, 2026, 2027, 2028, 2029])
calc.write_formula("B2", "=Inputs!B2", None, 1000)
for c in "CDEF":
    calc.write_formula(f"{c}2", f"={chr(ord(c) - 1)}2*(1+Growth)")
calc.add_sparkline("G2", {"range": "'Model Calc'!B2:F2"})
calc.add_table("A5:C8", {"name": "Tbl", "columns": [{"header": "K"}, {"header": "V"}, {"header": "Dbl", "formula": "=[@V]*2"}]})
for i in range(3):
    calc.write(5 + i, 0, f"k{i}"); calc.write(5 + i, 1, i + 1)
ch = wb.add_chart({"type": "line"}); ch.add_series({"values": "='Model Calc'!$B$2:$F$2"}); calc.insert_chart("A10", ch)
calc.write_dynamic_array_formula("H1", "=_xlfn._xlws.SORT(B1:F1,1,-1)")
calc.freeze_panes(1, 1)
wb.close()

def probe(p):
    z = zipfile.ZipFile(p); n = z.namelist()
    sx = "".join(z.read(k).decode("utf8", "ignore") for k in n if re.search(r"worksheets/sheet\d+\.xml$", k))
    styles = z.read("xl/styles.xml").decode()
    return {
        "chart": any("charts/chart" in k for k in n),
        "sparkline": "sparklineGroup" in sx,
        "table (ListObject)": any("tables/table" in k for k in n),
        "comment": any("comment" in k for k in n),
        "defined name": b"Growth" in z.read("xl/workbook.xml"),
        "data validation": "dataValidation" in sx,
        "conditional format": "conditionalFormatting" in sx,
        "freeze panes": "frozen" in sx,
        "dynamic array metadata": "xl/metadata.xml" in n,
        "cached formula values": bool(re.search(r"<f[^>]*>Inputs!B2</f><v>1000</v>", sx)),
        "custom number format kept verbatim": "$#,##0;($#,##0);-" in styles.replace("&quot;", ""),
    }

paths = {"original": SRC}
o = os.path.join(W, "openpyxl.xlsx"); openpyxl.load_workbook(SRC).save(o); paths["openpyxl load+save"] = o
paths["LibreOffice re-save"] = str(X.lo_recalc_copy(SRC))
p = xl.Patcher(SRC); p.apply([{"op": "set", "sheet": "Inputs", "cell": "B1", "value": 0.06}])
x = os.path.join(W, "xlpatch.xlsx"); p.pkg.save(x); paths["xl.py patch"] = x
res = {k: probe(v) for k, v in paths.items()}
feats = list(res["original"])
print("| feature | " + " | ".join(paths) + " |")
print("|---|" + "---|" * len(paths))
for f in feats:
    print(f"| {f} | " + " | ".join("kept" if res[k][f] else "LOST" for k in paths) + " |")
