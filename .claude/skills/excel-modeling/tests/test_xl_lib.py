"""Shared helpers for the xl.py tests."""
import re, zipfile, warnings
warnings.filterwarnings("ignore")
import xlcore as X
from openpyxl import load_workbook

def norm(f):
    """Normalise formula text; LibreOffice writes TRUE()/FALSE() for the constants."""
    if f is None:
        return None
    s = re.sub(r"\s+", "", str(f)).upper()
    return s.replace("FALSE()", "FALSE").replace("TRUE()", "TRUE")

def formulas(path):
    wb = load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and not re.fullmatch(r"=#[A-Z/0!?]+", c.value):
                    out[f"{ws.title}!{c.coordinate}"] = norm(c.value)   # LibreOffice turns cached errors into =#DIV/0!
        out[f"{ws.title}!CF"] = sorted(str(cf.sqref) for cf in ws.conditional_formatting)
        out[f"{ws.title}!DV"] = sorted(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    # x14 validations live in extLst, which openpyxl ignores; read them from the XML
    pkg = X.Package(path)
    for name, part, _ in pkg.sheets():
        x = pkg.raw[part].decode("utf8", "ignore")
        for block in re.findall(r"<x14:dataValidation .*?</x14:dataValidation>", x, re.S):
            for sq in re.findall(r"<xm:sqref>([^<]*)</xm:sqref>", block):
                out[f"{name}!DV"] = sorted(out.get(f"{name}!DV", []) + [sq])
    for n, d in wb.defined_names.items():
        out[f"name:{n}"] = norm(d.attr_text)
    return out

def chart_refs(path):
    z = zipfile.ZipFile(path)
    return sorted(norm(m) for n in z.namelist() if n.startswith("xl/charts/chart")
                  for m in re.findall(r"<c:f>([^<]+)</c:f>", z.read(n).decode())
                  if not m.startswith("&quot;"))  # LibreOffice writes series names as literal "<c:f>"name"

def values(path):
    rc = X.lo_recalc_copy(path)
    wb = load_workbook(rc, data_only=True)
    return {f"{ws.title}!{c.coordinate}": c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None}

